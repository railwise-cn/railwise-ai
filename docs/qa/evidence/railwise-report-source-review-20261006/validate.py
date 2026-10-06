from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import json, re, subprocess, zipfile
import pdfplumber
from PIL import Image, ImageDraw
from docx import Document
from openpyxl import load_workbook

root = Path(__file__).resolve().parent
renderer = '/Users/wangjiawei/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm'
render_dir = root / 'rendered'
render_dir.mkdir(exist_ok=True)
pdfs = sorted(root.glob('*.pdf'))
def render(pdf):
    subprocess.run([renderer, '-r', '120', '-png', str(pdf), str(render_dir / pdf.stem)], check=True, capture_output=True)
with ThreadPoolExecutor(max_workers=3) as pool:
    list(pool.map(render, pdfs))

results = {}
for pdf in pdfs:
    presentation = json.loads((root / f'{pdf.stem}.presentation.json').read_text())
    with pdfplumber.open(pdf) as document:
        pages = [page.extract_text() or '' for page in document.pages]
        text = '\n'.join(pages)
        (root / f'{pdf.stem}.text.txt').write_text(text)
        forbidden = [term for term in ['workwise-json-observation', 'adjustment_', 'cosa-in2-parser', 'P0 格式目录', 'network.instrumentParameters', 'SHA-256', 'inputHash', 'resultHash', 'projectionHash', 'normalized_data', '冻结的', '投影生成', '算法：', '算法 monitoring', 'deformation-internal-audit', 'monitoring-internal-audit'] if term in text]
        if forbidden:
            raise AssertionError(f'{pdf.name}: implementation terms {forbidden}')
        check = {'pages': len(pages), 'internalTermFindings': forbidden, 'headings': []}
        normalized_pages = [re.sub(r'\s+', '', page) for page in pages]
        for table in presentation['tables']:
            title = re.sub(r'\s+', '', table['title'])
            page_index = next((index for index, page in enumerate(normalized_pages) if title in page), None)
            if page_index is None:
                raise AssertionError(f'{pdf.name}: missing title {table["title"]}')
            first_row = table['rows'][0] if table['rows'] else None
            first_value = str(first_row.get(table['columns'][0]['key'], '不可用')) if first_row else '无可用记录/未评估'
            if re.sub(r'\s+', '', first_value) not in normalized_pages[page_index]:
                raise AssertionError(f'{pdf.name}: heading separated from first row {table["title"]}: {first_value}')
            check['headings'].append({'title': table['title'], 'page': page_index + 1, 'firstRowOnSamePage': True})
        if pdf.stem == 'long-table':
            for index in range(1, 146):
                assert f'观测-{index:03}' in text
            check['all145ObservationsPrinted'] = True
            check['continuationsRepeatHeader'] = all('观测号' in page for page in pages if re.search('观测-\d{3}', page))
            assert check['continuationsRepeatHeader']
        pngs = sorted(render_dir.glob(f'{pdf.stem}-*.png'), key=lambda path: int(path.stem.rsplit('-', 1)[1]))
        width, height = 850, 615
        contact = Image.new('RGB', (width * 2, height * ((len(pngs) + 1) // 2)), '#dce3ea')
        draw = ImageDraw.Draw(contact)
        for index, png in enumerate(pngs):
            image = Image.open(png).convert('RGB')
            image.thumbnail((width - 20, height - 36))
            x, y = (index % 2) * width + 10, (index // 2) * height + 28
            contact.paste(image, (x, y))
            draw.text((x, y - 20), f'{pdf.stem} / {index + 1}', fill='#17212b')
        contact.save(render_dir / f'{pdf.stem}-contact.png')
        results[pdf.name] = check

for stem in ['standard', 'segment', 'deformation', 'monitoring']:
    doc = Document(root / f'{stem}.docx')
    doc_text = '\n'.join([paragraph.text for paragraph in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    wb = load_workbook(root / f'{stem}.xlsx', data_only=True)
    wb_text = '\n'.join(str(cell.value or '') for sheet in wb for row in sheet for cell in row)
    for format, text in [('docx', doc_text), ('xlsx', wb_text)]:
        (root / f'{stem}.{format}.text.txt').write_text(text)
        for term in ['SHA-256', 'inputHash', 'resultHash', 'projectionHash', 'normalized_data', 'network.instrumentParameters', 'deformation-internal-audit', 'monitoring-internal-audit', '冻结的', '算法：']:
            assert term not in text, f'{stem}.{format}: {term}'
    results[f'{stem}.docx'] = {'tables': len(doc.tables), 'professionalTextParsed': True, 'internalTermFindings': []}
    results[f'{stem}.xlsx'] = {'sheets': len(wb.sheetnames), 'professionalTextParsed': True, 'internalTermFindings': [], 'numericCells': sum(isinstance(cell.value, (int, float)) for sheet in wb for row in sheet for cell in row)}

(root / 'validation-summary.json').write_text(json.dumps(results, ensure_ascii=False, indent=2))
print(json.dumps(results, ensure_ascii=False, indent=2))
